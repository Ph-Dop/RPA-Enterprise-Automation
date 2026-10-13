"""Tests for KPI computation and Excel report generation."""
from __future__ import annotations

from openpyxl import load_workbook

from rpa_platform.reporting.report_gen import ReportGenerator, compute_kpis


def test_compute_kpis(clean_df):
    kpis = compute_kpis(clean_df, currency="EUR")
    assert kpis["order_count"] == 3
    assert kpis["total_revenue"] == 120.0 + 540.5 + 210.0
    assert round(kpis["avg_order_value"], 2) == round((120.0 + 540.5 + 210.0) / 3, 2)
    assert kpis["currency"] == "EUR"
    assert kpis["monthly_revenue"] == {"2026-01": 120.0, "2026-02": 540.5, "2026-03": 210.0}


def test_empty_frame_kpis_are_zero():
    import pandas as pd

    kpis = compute_kpis(pd.DataFrame(columns=["amount", "ordered_at"]), currency="EUR")
    assert kpis["total_revenue"] == 0.0
    assert kpis["order_count"] == 0
    assert kpis["avg_order_value"] == 0.0


def test_export_creates_workbook_with_expected_sheets(settings, clean_df):
    generator = ReportGenerator(settings)
    quality = {
        "raw_count": 3,
        "cleaned_count": 3,
        "removed_total": 0,
        "issue_counts": {"missing_fields": 0, "duplicates": 0, "invalid_amount": 0, "invalid_date": 0},
        "issues": {},
        "total_revenue_after": 870.5,
        "order_count_after": 3,
    }
    path = generator.export(clean_df, quality)

    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Quality Report", "Order Details"]

    summary = wb["Summary"]
    assert "Test GmbH" in summary["A1"].value
    assert "870.50" in str(summary["B3"].value)  # total revenue value
