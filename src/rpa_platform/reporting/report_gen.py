"""Automatic business report generation.

Computes KPIs (total revenue, order count, average order value, monthly
breakdown) and exports a formatted Excel workbook:
- Sheet 1 "Summary"      - KPI dashboard with styled cells
- Sheet 2 "Quality Report" - cleaning statistics (before/after, issue classes)
- Sheet 3 "Order Details"  - cleaned orders

Files are named ``orders_report_YYYYMMDD_HHMMSS.xlsx`` and stored in the
configured output directory, ready to be scheduled weekly/monthly.
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd

from ..config import Settings

logger = logging.getLogger("rpa.reporting")


def compute_kpis(df: pd.DataFrame, currency: str = "EUR") -> dict:
    """Compute headline KPIs from cleaned order data."""
    total_revenue = float(df["amount"].sum()) if not df.empty else 0.0
    order_count = int(len(df))
    avg_order_value = total_revenue / order_count if order_count else 0.0

    monthly: dict[str, float] = {}
    if not df.empty and "ordered_at" in df.columns:
        months = df["ordered_at"].astype(str).str[:7]
        monthly = {k: round(float(v), 2) for k, v in df.groupby(months)["amount"].sum().items()}

    return {
        "total_revenue": round(total_revenue, 2),
        "order_count": order_count,
        "avg_order_value": round(avg_order_value, 2),
        "currency": currency,
        "monthly_revenue": dict(sorted(monthly.items())),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }


class ReportGenerator:
    """Writes the formatted Excel business report."""

    HEADER_FILL = "1F4E79"
    HEADER_FONT_COLOR = "FFFFFF"
    KPI_FILL = "DDEBF7"

    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.output_dir: Path = settings.report_output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        cleaned: pd.DataFrame,
        quality_report: dict,
        kpis: dict | None = None,
    ) -> Path:
        """Generate the report workbook and return the output file path."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        kpis = kpis or compute_kpis(cleaned, self.settings.report_currency_symbol)

        wb = Workbook()

        # ---- Sheet 1: Summary ------------------------------------------
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = f"Business Report - {self.settings.company_name}"
        ws["A1"].font = Font(size=16, bold=True, color=self.HEADER_FONT_COLOR)
        ws["A1"].fill = PatternFill("solid", fgColor=self.HEADER_FILL)
        ws.merge_cells("A1:C1")

        kpi_rows = [
            ("Total Revenue", f"{kpis['currency']} {kpis['total_revenue']:,.2f}"),
            ("Order Count", f"{kpis['order_count']:,}"),
            ("Average Order Value", f"{kpis['currency']} {kpis['avg_order_value']:,.2f}"),
        ]
        for idx, (label, value) in enumerate(kpi_rows, start=3):
            ws.cell(row=idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=idx, column=1).fill = PatternFill("solid", fgColor=self.KPI_FILL)
            ws.cell(row=idx, column=2, value=value)
            ws.cell(row=idx, column=2).number_format = "0.00"

        ws.cell(row=8, column=1, value="Monthly Revenue").font = Font(bold=True)
        month_start = 9
        for idx, (month, amount) in enumerate(kpis["monthly_revenue"].items()):
            ws.cell(row=month_start + idx, column=1, value=str(month))
            cell = ws.cell(row=month_start + idx, column=2, value=float(amount))
            cell.number_format = "#,##0.00"

        gen_row = month_start + len(kpis["monthly_revenue"]) + 2
        ws.cell(row=gen_row, column=1, value=f"Generated: {kpis['generated_at']}")

        # ---- Sheet 2: Quality Report -------------------------------------
        ws2 = wb.create_sheet("Quality Report")
        ws2["A1"] = "Data Quality Report"
        ws2["A1"].font = Font(size=14, bold=True, color=self.HEADER_FONT_COLOR)
        ws2["A1"].fill = PatternFill("solid", fgColor=self.HEADER_FILL)
        ws2.merge_cells("A1:C1")

        quality_rows = [
            ("Raw orders received", quality_report["raw_count"]),
            ("Valid orders after cleaning", quality_report["cleaned_count"]),
            ("Removed (total)", quality_report["removed_total"]),
        ]
        for i, (label, value) in enumerate(quality_rows, start=3):
            ws2.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws2.cell(row=i, column=2, value=value)

        ws2.cell(row=7, column=1, value="Issue Class").font = Font(bold=True)
        ws2.cell(row=7, column=2, value="Count").font = Font(bold=True)
        issue_counts = quality_report.get("issue_counts", {})
        for i, (cls, count) in enumerate(issue_counts.items(), start=8):
            ws2.cell(row=i, column=1, value=cls)
            ws2.cell(row=i, column=2, value=count)

        # ---- Sheet 3: Order Details --------------------------------------
        ws3 = wb.create_sheet("Order Details")
        if cleaned.empty:
            ws3["A1"] = "No valid orders in this period."
        else:
            for col_idx, col_name in enumerate(cleaned.columns, start=1):
                header = ws3.cell(row=1, column=col_idx, value=col_name)
                header.font = Font(bold=True, color=self.HEADER_FONT_COLOR)
                header.fill = PatternFill("solid", fgColor=self.HEADER_FILL)
            for row_idx, record in enumerate(cleaned.itertuples(index=False), start=2):
                for col_idx, value in enumerate(record, start=1):
                    cell = ws3.cell(row=row_idx, column=col_idx, value=value)
                    if col_idx == 5:  # amount column
                        cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(vertical="top")

        # ---- column widths -------------------------------------------------
        for sheet in (ws, ws2, ws3):
            for col in range(1, 4):
                sheet.column_dimensions[get_column_letter(col)].width = 28

        out_path = self.output_dir / f"orders_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        wb.save(out_path)
        logger.info("Report saved to %s", out_path)
        return out_path
